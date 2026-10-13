"""
reports.py
==========
Excel reports (Phase 5.9). Every workbook has the same layout: a title block (report, filters, generated time, data-quality line),
a styled header row with filters and frozen panes, typed cells with number formats, and a Notes sheet that states the definitions
and the caveats — in particular that ACTUAL penalties (printed on bills) and ESTIMATED savings (theoretical) are never combined.
"""

import io
import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import analytics, settings
from .analytics import FROM, Filters

NAVY = "082B49"
INR = '#,##0.00;[Red]-#,##0.00'
NUM = '#,##0.00'
INT = '#,##0'
PF = '0.0000'
PCT = '0.0"%"'

REGISTER_SELECT = """SELECT b.id AS bill_id, b.source, b.plant_code, pl.plant_name, b.state, b.discom_code, b.consumer_number, b.meter_number, b.po_number,
    b.bill_number, b.bill_date, b.due_date, b.period_start, b.period_end, b.billing_month, b.billing_days, b.billing_basis, b.kwh, b.kvah,
    b.pf_billed, b.pf_calculated, b.sanctioned_load, b.sanctioned_load_unit, b.contract_demand, b.contract_demand_unit, b.max_demand, b.max_demand_unit,
    b.billing_demand, b.billing_demand_unit, b.energy_charges, b.fixed_charges, b.demand_charges, b.fuel_adjustment, b.electricity_duty, b.wheeling,
    b.taxes_surcharges, b.arrears, b.adjustments_credits, b.late_payment, b.pf_penalty, b.pf_penalty_status, b.pf_incentive, b.pf_incentive_status,
    b.excess_demand_penalty, b.excess_demand_status, b.total_bill_amount, b.total_payable, b.cost_amount, b.cost_source, b.data_flags, b.bill_status,
    b.quality, b.is_primary, b.revision_count,
    pf.pf_difference, pf.validation_status AS pf_validation, pf.pf_used, pf.pf_source, pf.performance_status, pf.direction,
    pf.expected_pf_penalty, pf.expected_status AS pf_expected_status, pf.tariff_rule_id AS pf_rule, pf.current_kvah, pf.target_pf, pf.target_kvah,
    pf.kvah_reduction, pf.marginal_rate, pf.rate_source, pf.estimated_savings, pf.savings_status,
    d.comparison_status AS demand_comparison, d.utilization_pct, d.excess_demand, d.exceeds_contract, d.violation_status, d.tolerance_pct,
    d.tariff_rule_id AS demand_rule, d.expected_excess_penalty, d.expected_status AS demand_expected_status"""

C = {  # key -> (header, number format, width)
    "plant_code": ("Plant code", None, 11), "plant_name": ("Plant", None, 28), "state": ("State", None, 16), "discom_code": ("Electricity board", None, 16),
    "consumer_number": ("Consumer no.", None, 17), "meter_number": ("Meter no.", None, 14), "po_number": ("PO no.", None, 12), "bill_number": ("Bill no.", None, 16),
    "bill_date": ("Bill date", None, 11), "due_date": ("Due date", None, 11), "period_start": ("Period start", None, 11), "period_end": ("Period end", None, 11),
    "billing_month": ("Billing month", None, 10), "billing_days": ("Days", INT, 6), "billing_basis": ("Billing basis", None, 9), "kwh": ("kWh", NUM, 13),
    "kvah": ("kVAh", NUM, 13), "pf_billed": ("PF as per bill", PF, 10), "pf_calculated": ("PF calculated (kWh/kVAh)", PF, 12), "pf_difference": ("Difference", PF, 10),
    "pf_validation": ("PF validation", None, 24), "pf_used": ("PF used", PF, 9), "pf_source": ("PF source", None, 22), "performance_status": ("PF status", None, 11),
    "direction": ("Lead/Lag", None, 9), "contract_demand": ("Contract demand", NUM, 11), "contract_demand_unit": ("CD unit", None, 8),
    "max_demand": ("Maximum demand", NUM, 11), "max_demand_unit": ("MD unit", None, 8), "billing_demand": ("Billing demand", NUM, 11),
    "billing_demand_unit": ("BD unit", None, 8), "sanctioned_load": ("Sanctioned load", NUM, 11), "sanctioned_load_unit": ("SL unit", None, 8),
    "utilization_pct": ("Demand utilisation %", PCT, 11), "excess_demand": ("Excess over contract demand", NUM, 12), "violation_status": ("Violation status", None, 24),
    "tolerance_pct": ("Permitted tolerance %", NUM, 10), "demand_rule": ("Tariff rule #", INT, 8), "pf_rule": ("Tariff rule #", INT, 8),
    "energy_charges": ("Energy charges", INR, 14), "fixed_charges": ("Fixed charges", INR, 14), "demand_charges": ("Demand charges", INR, 14),
    "fuel_adjustment": ("Fuel adjustment", INR, 13), "electricity_duty": ("Electricity duty", INR, 13), "wheeling": ("Wheeling", INR, 12),
    "taxes_surcharges": ("Taxes & surcharges", INR, 14), "arrears": ("Arrears", INR, 13), "adjustments_credits": ("Adjustments / credits", INR, 14),
    "late_payment": ("Late payment", INR, 12), "pf_penalty": ("PF penalty (actual, on bill)", INR, 15), "pf_penalty_status": ("PF penalty status", None, 12),
    "pf_incentive": ("PF incentive (actual, on bill)", INR, 15), "pf_incentive_status": ("PF incentive status", None, 12),
    "excess_demand_penalty": ("Excess-demand penalty (actual, on bill)", INR, 17), "excess_demand_status": ("Excess-demand status", None, 13),
    "total_bill_amount": ("Total bill amount", INR, 15), "total_payable": ("Total payable", INR, 15), "cost_amount": ("Cost used in totals", INR, 15),
    "cost_source": ("Cost source", None, 11), "data_flags": ("Consistency flags", None, 24), "bill_status": ("Bill status", None, 13), "quality": ("Verification", None, 12),
    "source": ("Source", None, 14), "revision_count": ("Bills for this month", INT, 9), "is_primary": ("Counted in totals", None, 9),
    "expected_pf_penalty": ("PF penalty expected by approved tariff rule", INR, 18), "pf_expected_status": ("Expected-penalty status", None, 26),
    "expected_excess_penalty": ("Excess-demand penalty expected by approved tariff rule", INR, 18), "demand_expected_status": ("Expected-excess status", None, 26),
    "current_kvah": ("Current kVAh (kWh/PF)", NUM, 13), "target_pf": ("Target PF", PF, 9), "target_kvah": ("Target kVAh (kWh/target PF)", NUM, 14),
    "kvah_reduction": ("Potential kVAh reduction", NUM, 13), "marginal_rate": ("Rate used (Rs/kVAh)", NUM, 11), "rate_source": ("Rate source", None, 24),
    "estimated_savings": ("Estimated saving (THEORETICAL, not billed)", INR, 20), "savings_status": ("Saving status", None, 26),
    "demand_comparison": ("Demand comparison", None, 18), "exceeds_contract": ("MD above CD", None, 8),
}

REPORTS = {
    "monitor": ("Plant-wise Electricity Monitoring", ["plant_code", "plant_name", "state", "discom_code", "period_start", "period_end", "billing_basis", "kwh", "kvah", "pf_used",
                "performance_status", "contract_demand", "contract_demand_unit", "max_demand", "max_demand_unit", "pf_penalty", "excess_demand_penalty", "total_payable",
                "bill_status", "quality"], None),
    "consumption": ("Plant-wise Electricity Consumption", ["plant_code", "plant_name", "state", "discom_code", "consumer_number", "meter_number", "billing_month", "period_start",
                    "period_end", "billing_days", "billing_basis", "kwh", "kvah", "pf_billed", "cost_amount", "bill_status", "quality"], None),
    "pf": ("PF Performance", ["plant_code", "plant_name", "state", "discom_code", "billing_month", "billing_basis", "kwh", "kvah", "pf_billed", "pf_calculated", "pf_difference",
           "pf_validation", "pf_used", "pf_source", "performance_status", "direction", "bill_status", "quality"], None),
    "pf_penalties": ("Actual PF Penalties", ["plant_code", "plant_name", "state", "discom_code", "billing_month", "bill_number", "po_number", "pf_billed", "billing_basis",
                     "pf_penalty", "pf_penalty_status", "pf_incentive", "expected_pf_penalty", "pf_expected_status", "pf_rule", "bill_status", "quality"],
                     "(b.pf_penalty > 0 OR b.pf_penalty_status = 'UNREADABLE' OR b.pf_incentive > 0)"),
    "demand_violations": ("Demand Violations", ["plant_code", "plant_name", "state", "discom_code", "billing_month", "bill_number", "contract_demand", "contract_demand_unit",
                          "max_demand", "max_demand_unit", "billing_demand", "utilization_pct", "excess_demand", "tolerance_pct", "violation_status", "demand_rule",
                          "excess_demand_penalty", "demand_comparison", "bill_status", "quality"],
                          "(d.exceeds_contract = 1 OR d.violation_status IN ('VIOLATION','VIOLATION_BILLED'))"),
    "demand_penalties": ("Actual Demand Penalties", ["plant_code", "plant_name", "state", "discom_code", "billing_month", "bill_number", "po_number", "contract_demand",
                         "contract_demand_unit", "max_demand", "max_demand_unit", "demand_charges", "excess_demand_penalty", "excess_demand_status", "expected_excess_penalty",
                         "demand_expected_status", "demand_rule", "bill_status", "quality"],
                         "(b.excess_demand_penalty > 0 OR b.excess_demand_status = 'UNREADABLE')"),
    "monthly_cost": ("Monthly Electricity Cost", ["plant_code", "plant_name", "state", "discom_code", "billing_month", "kwh", "kvah", "energy_charges", "fixed_charges",
                     "demand_charges", "fuel_adjustment", "electricity_duty", "wheeling", "taxes_surcharges", "arrears", "adjustments_credits", "total_bill_amount",
                     "total_payable", "cost_amount", "cost_source", "data_flags", "bill_status", "quality"], None),
    "savings": ("Estimated Savings (theoretical)", ["plant_code", "plant_name", "state", "discom_code", "billing_month", "billing_basis", "kwh", "pf_used", "target_pf",
                "current_kvah", "target_kvah", "kvah_reduction", "marginal_rate", "rate_source", "estimated_savings", "savings_status", "quality"],
                "pf.savings_status LIKE 'ESTIMATE%'"),
    "register": ("Full Electricity Bill Register", list(C) , None),
}

NOTES = [
    "ACTUAL penalties are amounts printed on the electricity bills. ESTIMATED savings are theoretical kVAh reductions if PF were raised to the target PF. "
    "They are different things and are never added together.",
    "Every figure was read from a bill by software. A bill whose 'Verification' is UNVERIFIED has not been checked by a person.",
    "Blank cells mean 'not found / not applicable / not readable' - they are not zero. A zero means the bill printed zero.",
    "Only the latest bill of an account and month is counted in totals (revised or duplicate bills are kept in the Full Register, 'Counted in totals' = No).",
    "Energy-weighted PF = total kWh / total kVAh over bills where both are plausible and agree with the PF printed on the bill. It is not the average of PFs.",
    "A bill with contradictory figures (see 'Consistency flags') stays in the report but its inconsistent figure is left out of company totals.",
    "A PF penalty / demand violation expected by an APPROVED tariff rule is shown only where such a rule exists; otherwise the status says UNAVAILABLE_NO_RULE.",
    "Recorded demand above contract demand is NOT automatically a penalty or a violation: that depends on the tariff's permitted limit.",
]


def _scope_text(f: Filters) -> str:
    parts = []
    for label, v in (("plants", f.plants), ("states", f.states), ("boards", f.discoms)):
        if v:
            parts.append(f"{label}: {', '.join(v[:6])}{'...' if len(v) > 6 else ''}")
    if f.month_from or f.month_to:
        parts.append(f"billing months {f.month_from or '...'} to {f.month_to or '...'}")
    for label, v in (("basis", f.basis), ("PF status", f.pf_status), ("exception", f.exception), ("search", f.search)):
        if v:
            parts.append(f"{label}: {v}")
    parts.append("verified bills only" if f.quality == "verified" else "all machine-read bills")
    return "; ".join(parts) or "all data"


def _sheet(ws, title, headers, widths, rows, fmts, f: Filters, conn, kpi_line=None):
    ws.title = "Report"
    ws["A1"] = title
    ws["A1"].font = Font(name="Calibri", size=15, bold=True, color=NAVY)
    ws["A2"] = f"Scope: {_scope_text(f)}"
    ws["A3"] = f"Generated {datetime.now():%d %b %Y %H:%M} | RDC Electricity Intelligence"
    ws["A4"] = kpi_line or ""
    for r in (2, 3, 4):
        ws[f"A{r}"].font = Font(size=9, italic=True, color="555555")
    hr = 6
    thin = Side(style="thin", color="C9D1DA")
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=hr, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        ws.column_dimensions[get_column_letter(j)].width = widths[j - 1]
    ws.row_dimensions[hr].height = 44
    for i, row in enumerate(rows, start=hr + 1):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=v)
            if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
                c.data_type = "s"                      # text read from a bill must never be executed as a spreadsheet formula
            if fmts[j - 1] and isinstance(v, (int, float)):
                c.number_format = fmts[j - 1]
        # (no per-cell border: building a style object for every one of 100k+ cells made big exports take a minute)
    ws.freeze_panes = ws.cell(row=hr + 1, column=3)
    if rows:
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(headers))}{hr + len(rows)}"
    return hr + len(rows)


def _notes_sheet(wb, extra=()):
    ws = wb.create_sheet("Notes")
    ws.column_dimensions["A"].width = 130
    ws["A1"] = "Definitions and caveats"
    ws["A1"].font = Font(size=13, bold=True, color=NAVY)
    for i, t in enumerate(list(NOTES) + list(extra), start=3):
        ws.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True, vertical="top")


def _dq_line(conn, f: Filters) -> str:
    k = analytics.kpis(conn, f)
    return (f"{k['bills']} bills in scope, {k['unverified']} not yet verified by a person; "
            f"{k['pairs_excluded']} kWh/kVAh pairs and {k['cost_excluded']} costs left out of totals because the bill contradicts itself.")


def build(conn, kind: str, f: Filters):
    """(filename, xlsx bytes) for one of REPORTS or 'states' / 'exceptions'."""
    wb = Workbook()
    ws = wb.active
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    if kind in REPORTS:
        title, cols, extra_where = REPORTS[kind]
        f2 = Filters(**{**f.__dict__, "primary_only": kind != "register"})
        w, p = f2.where()
        sql = f"{REGISTER_SELECT} {FROM} WHERE {w}" + (f" AND {extra_where}" if extra_where else "") + \
              " ORDER BY b.plant_code, b.billing_month, b.id"
        data = [dict(r) for r in conn.execute(sql, p)]
        rows = [[(("Yes" if r[c] else "No") if c in ("is_primary", "exceeds_contract") and r[c] is not None else r[c]) for c in cols] for r in data]
        last = _sheet(ws, title, [C[c][0] for c in cols], [C[c][2] for c in cols], rows, [C[c][1] for c in cols], f, conn, _dq_line(conn, f2))
        # totals row for the money columns that are ACTUAL amounts
        if kind in ("pf_penalties", "demand_penalties"):
            col = "pf_penalty" if kind == "pf_penalties" else "excess_demand_penalty"
            j = cols.index(col) + 1
            ws.cell(row=last + 2, column=1, value="Total of actual amounts printed on the bills (unreadable lines are not included)").font = Font(bold=True)
            tc = ws.cell(row=last + 2, column=j, value=sum(r[col] or 0 for r in data if (r[col] or 0) > 0))
            tc.number_format, tc.font = INR, Font(bold=True)
        if kind == "savings":
            j = cols.index("estimated_savings") + 1
            ws.cell(row=last + 2, column=1, value="Total THEORETICAL saving (not a billed amount, not a penalty)").font = Font(bold=True)
            tc = ws.cell(row=last + 2, column=j, value=sum(r["estimated_savings"] or 0 for r in data))
            tc.number_format, tc.font = INR, Font(bold=True)
        _notes_sheet(wb)
        name = f"RDC_EB_{kind}_{stamp}.xlsx"
    elif kind == "states":
        rows_ = analytics.group_performance(conn, f, "state")
        cols = [("State", None, 20), ("Bills", INT, 8), ("Plants", INT, 8), ("PF (energy-weighted)", PF, 12), ("Actual PF penalties", INR, 15),
                ("Actual excess-demand penalties", INR, 17), ("Demand violations", INT, 10), ("Cost", INR, 16), ("Theoretical savings", INR, 16), ("Unverified bills", INT, 9)]
        rows = [[r["grp"], r["bills"], r["plants"], r["pf"], r["pf_penalty"], r["xd_penalty"], r["violations"], r["cost"], r["savings"], r["unverified"]] for r in rows_]
        _sheet(ws, "State-wise Summary", [c[0] for c in cols], [c[2] for c in cols], rows, [c[1] for c in cols], f, conn, _dq_line(conn, f))
        boards = wb.create_sheet("By electricity board")
        rows_b = analytics.group_performance(conn, f, "discom")
        _sheet(boards, "Electricity-board-wise Summary", ["Board"] + [c[0] for c in cols[1:]], [c[2] for c in cols],
               [[r["grp"], r["bills"], r["plants"], r["pf"], r["pf_penalty"], r["xd_penalty"], r["violations"], r["cost"], r["savings"], r["unverified"]] for r in rows_b],
               [c[1] for c in cols], f, conn)
        boards.title = "By electricity board"
        _notes_sheet(wb)
        name = f"RDC_EB_state_summary_{stamp}.xlsx"
    elif kind == "exceptions":
        name = f"RDC_EB_exceptions_{stamp}.xlsx"
        allowed = f.allowed_plants
        scope, p = "", {}
        if allowed is not None:
            keys = [f":a{i}" for i in range(len(allowed))] or ["NULL"]
            p = {f"a{i}": v for i, v in enumerate(allowed)}
            scope = f"AND COALESCE(pp.plant_code, u.plant_code) IN ({','.join(keys)})"
        rows_ = conn.execute(f"""SELECT e.id, e.bill_status, COALESCE(pp.plant_code, u.plant_code) AS plant, COALESCE(pp.po_number, u.po_number) AS po, e.discom_code,
            e.state, e.doc_kind, e.pages, e.ocr_pages, e.created_at, e.missing_json, e.flags_json, e.warnings_json, e.error, e.needs_review_count,
            CASE WHEN e.attachment_id IS NOT NULL THEN 'ORACLE_EBS' ELSE 'MANUAL_UPLOAD' END AS source,
            (SELECT GROUP_CONCAT(field_name || ' (' || ROUND(COALESCE(confidence, 0), 2) || ')') FROM electricity_bill_fields f
              WHERE f.extraction_id = e.id AND f.verification_status = 'NEEDS_REVIEW') AS low_conf
            FROM electricity_bill_extractions e LEFT JOIN electricity_attachments a ON a.attached_document_id = e.attachment_id
            LEFT JOIN electricity_pos pp ON pp.po_header_id = a.po_header_id LEFT JOIN electricity_uploads u ON u.id = e.upload_id
            WHERE e.archived = 0 AND e.bill_status <> 'VERIFIED' {scope} ORDER BY e.bill_status, plant, e.id""", p).fetchall()
        hdr = [("Extraction", INT, 10), ("Status", None, 14), ("Plant", None, 10), ("PO", None, 12), ("Board", None, 14), ("State", None, 14), ("Source", None, 13),
               ("Document", None, 10), ("Pages", INT, 7), ("OCR pages", INT, 8), ("Read on", None, 17), ("Missing fields", None, 36), ("Low-confidence fields (confidence)", None, 46),
               ("Flags / problems", None, 70), ("Error", None, 40)]
        rows = []
        for r in rows_:
            flags = "; ".join(json.loads(r["flags_json"] or "[]") + json.loads(r["warnings_json"] or "[]"))
            rows.append([r["id"], r["bill_status"], r["plant"], r["po"], r["discom_code"], r["state"], r["source"], r["doc_kind"], r["pages"], r["ocr_pages"], r["created_at"],
                         ", ".join(json.loads(r["missing_json"] or "[]")), r["low_conf"], flags, r["error"]])
        _sheet(ws, "Exception / Failed Bill Report", [h[0] for h in hdr], [h[2] for h in hdr], rows, [h[1] for h in hdr], f, conn,
               f"{len(rows)} extractions are not verified (needs review, partial, failed, unsupported, not a bill or duplicate).")
        _notes_sheet(wb)
    else:
        raise ValueError(f"unknown report '{kind}'")
    buf = io.BytesIO()
    wb.save(buf)
    return name, buf.getvalue()


ALL_KINDS = [(k, v[0]) for k, v in REPORTS.items()] + [("states", "State-wise Summary"), ("exceptions", "Exception / Failed Bill Report")]
