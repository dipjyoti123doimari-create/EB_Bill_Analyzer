"""Phase 5 tests: the website — login and roles, plant scoping, CSRF, every page, bill viewer, corrections, alerts, tariff rules, settings,
manual upload through the browser, Excel reports, and the security properties (no secrets in pages, escaping, SQL-injection, path safety).
"""
import io
import json
import os
import pathlib
import re
import sqlite3
import sys
import tempfile
import unittest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))

from openpyxl import load_workbook
from werkzeug.security import generate_password_hash

import helpers
import synth
from eb_app import auth, config, jobs, storage
from eb_app.reader import store

ORACLE_SECRET = "OracleSecret#12345"
SMTP_SECRET = "MailSecret#67890"


class WebBase(unittest.TestCase):
    """A website over a temporary database with a few real (synthetic) bills, and RDC-OPS-style users in a temporary users database."""

    @classmethod
    def build(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        t = pathlib.Path(cls.tmp.name)
        cls._orig = (config.DATA_DIR, config.STORAGE_DIR, config.ORIGINALS_DIR, os.environ.get("EB_AUTH_DB"), os.environ.get("EB_ORACLE_PASSWORD"), os.environ.get("EB_SMTP_PASSWORD"))
        config.DATA_DIR, config.STORAGE_DIR = t, t / "storage"
        config.ORIGINALS_DIR = config.STORAGE_DIR / "originals"
        cls.ops = t / "ops.db"
        c = sqlite3.connect(cls.ops)
        c.executescript("create table users (id integer primary key, full_name text, email text, username text, password_hash text, role text, is_active integer);"
                        "create table user_plant_access (id integer primary key, user_id integer, plant_code text, plant_name text);")
        for i, (u, role, active) in enumerate([("admin", "SUPER_ADMIN", 1), ("viewer", "HO_VIEWER", 1), ("plant1", "PLANT_USER", 1), ("uep", "UEP_ADMIN", 1),
                                               ("gone", "SUPER_ADMIN", 0), ("plant2", "PLANT_USER", 1), ("noplants", "REGIONAL_USER", 1)], start=1):
            c.execute("insert into users values (?,?,?,?,?,?,?)", (i, u.title(), f"{u}@x.in", u, generate_password_hash("pw-" + u), role, active))
        c.execute("insert into user_plant_access values (1, 3, 'PU1', 'Pune-1')")
        c.execute("insert into user_plant_access values (2, 6, 'NA1', 'Nashik-1')")
        c.commit()
        c.close()
        os.environ["EB_AUTH_DB"] = str(cls.ops)
        os.environ["EB_ORACLE_PASSWORD"] = ORACLE_SECRET
        os.environ["EB_SMTP_PASSWORD"] = SMTP_SECRET
        cls.db = t / "eb.db"
        storage.init_db(cls.db)
        with storage.session(cls.db) as conn:
            for pid, plant, name, sup in ((1, "PU1", "Pune-1", "MSEDCL PUNE <script>alert(1)</script>"), (2, "NA1", "Nashik-1", "MSEDCL NASHIK"), (3, "KT1", "Kottayam", "KSEB KOTTAYAM")):
                conn.execute("INSERT INTO electricity_pos (po_header_id, po_number, supplier_name, plant_code, plant_name, billing_start, billing_end, first_seen, last_seen)"
                             " VALUES (?,?,?,?,?,?,?,?,?)", (pid, f"P{pid}", sup, plant, name, "2026-09-01", "2026-09-30", "x", "x"))
        cls.exts = {}
        for att, po, key, kw in ((501, 1, "pu", dict(bill_no="W-0001", pf_penalty=5000.0)), (502, 2, "na", dict(bill_no="W-0002", consumer="222222222222", kwh=60000, kvah=100000, pf=0.60)),
                                 (503, 3, "kt", dict(bill_no="W-0003", consumer="333333333333", discom_line="Kerala State Electricity Board Ltd (KSEB)"))):
            path = config.ORIGINALS_DIR / "t" / f"w{att}.pdf"
            synth.make_bill(path, **kw)
            with storage.session(cls.db) as conn:
                conn.execute("INSERT INTO electricity_attachments (attached_document_id, po_header_id, entity_name, pk1_value, document_id, file_name, fetch_status, sha256, local_path,"
                             " first_seen, last_seen) VALUES (?,?,?,?,?,?,?,?,?,?,?)", (att, po, "PO_HEADERS", str(po), att, f"w{att}.pdf", "DOWNLOADED", f"sha{att}", f"originals/t/w{att}.pdf", "x", "x"))
            cls.exts[key] = store.process_attachment(att, db_path=cls.db, run_ocr=False)["extraction_id"]
        import app as appmod
        cls.app = appmod.create_app(db_path=cls.db, start_background=False, testing=True)

    @classmethod
    def destroy(cls):
        config.DATA_DIR, config.STORAGE_DIR, config.ORIGINALS_DIR = cls._orig[:3]
        for k, v in zip(("EB_AUTH_DB", "EB_ORACLE_PASSWORD", "EB_SMTP_PASSWORD"), cls._orig[3:]):
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
        cls.tmp.cleanup()

    def login(self, user, password=None, client=None):
        auth.reset_throttle()
        c = client or self.app.test_client()
        r = c.post("/login", data={"username": user, "password": password if password is not None else "pw-" + user})
        self.assertEqual(r.status_code, 302, f"login of {user} failed")
        return c

    def token(self, c):
        with c.session_transaction() as s:
            return s["csrf"]

    def post(self, c, url, data=None, **kw):
        d = dict(data or {})
        return c.post(url, data=d, headers={"X-CSRF-Token": self.token(c)}, **kw)

    def rows(self, sql, *a):
        with storage.session(self.db) as conn:
            return [dict(r) for r in conn.execute(sql, a).fetchall()]


# ------------------------------------------------------------------------------------------------ read-only behaviour
class ReadOnlyWebTests(WebBase):
    @classmethod
    def setUpClass(cls):
        cls.build()

    @classmethod
    def tearDownClass(cls):
        cls.destroy()

    PAGES = ["/", "/?period=last3", "/monitor", "/monitor?sort=pf&dir=asc", "/monitor?q=PU1&state=Maharashtra&discom=MSEDCL&plant=PU1&basis=Unknown", "/monitor?year=2026&month=9",
             "/monitor?exception=pf_penalty", "/monitor?pf_status=CRITICAL", "/plant/PU1", "/review", "/review?status=ALL", "/alerts", "/alerts?status=ALL", "/upload", "/reports",
             "/automation", "/api/automation", "/tariffs", "/settings", "/rules", "/discovery", "/audit", "/api/dashboard", "/api/oracle-status", "/healthz"]

    def test_every_page_renders_for_an_administrator(self):
        c = self.login("admin")
        for url in self.PAGES:
            r = c.get(url)
            self.assertEqual(r.status_code, 200, url)

    def test_anonymous_visitors_are_sent_to_the_login_page_and_apis_answer_401(self):
        c = self.app.test_client()
        for url in ("/", "/monitor", "/plant/PU1", "/review", "/alerts", "/upload", "/reports", "/automation", "/settings", f"/e/{self.exts['pu']}", f"/e/{self.exts['pu']}/file"):
            r = c.get(url)
            self.assertEqual(r.status_code, 302, url)
            self.assertIn("/login", r.headers["Location"])
        for url in ("/api/dashboard", "/api/automation", "/api/upload/1", "/api/oracle-status"):
            self.assertEqual(c.get(url, headers={"Accept": "application/json"}).status_code, 401, url)
        self.assertEqual(c.get("/reports/pf.xlsx").status_code, 302)
        self.assertEqual(c.get("/healthz").status_code, 200)                       # a watchdog needs this without a login
        self.assertEqual(c.get("/login").status_code, 200)

    def test_login_failures_say_the_same_thing_and_a_disabled_or_foreign_role_cannot_enter(self):
        auth.reset_throttle()
        c = self.app.test_client()
        msgs = set()
        for u, p in (("admin", "wrong"), ("nobody", "x"), ("gone", "pw-gone")):
            r = c.post("/login", data={"username": u, "password": p})
            self.assertEqual(r.status_code, 401)
            msgs.add(re.search(r'alert-error">([^<]+)<', r.get_data(as_text=True)).group(1))
        self.assertEqual(msgs, {"Invalid username or password."})                   # no account enumeration
        r = c.post("/login", data={"username": "uep", "password": "pw-uep"})
        self.assertEqual(r.status_code, 401)
        self.assertIn("no access", r.get_data(as_text=True))

    def test_repeated_wrong_passwords_lock_that_user_and_address_out(self):
        auth.reset_throttle()
        c = self.app.test_client()
        for _ in range(auth.MAX_FAILURES):
            c.post("/login", data={"username": "viewer", "password": "bad"})
        r = c.post("/login", data={"username": "viewer", "password": "pw-viewer"})           # even the right password is refused for now
        self.assertEqual(r.status_code, 401)
        self.assertIn("Too many", r.get_data(as_text=True))
        auth.reset_throttle()
        self.assertEqual(c.post("/login", data={"username": "viewer", "password": "pw-viewer"}).status_code, 302)

    def test_session_cookie_has_its_own_name_and_is_http_only(self):
        c = self.app.test_client()
        auth.reset_throttle()
        r = c.post("/login", data={"username": "admin", "password": "pw-admin"})
        cookie = r.headers.get("Set-Cookie", "")
        self.assertIn("eb_session=", cookie)                                           # never collides with the RDC-OPS session cookie
        self.assertIn("HttpOnly", cookie)
        self.assertIn("SameSite=Lax", cookie)
        self.assertEqual(r.headers.get("X-Content-Type-Options"), "nosniff")

    def test_role_matrix(self):
        admin, viewer, plant = self.login("admin"), self.login("viewer"), self.login("plant1")
        admin_only = ["/automation", "/settings", "/tariffs", "/rules", "/discovery", "/audit", "/api/automation"]
        for url in admin_only:
            self.assertEqual(admin.get(url).status_code, 200, url)
            self.assertEqual(viewer.get(url).status_code, 403, url)
            self.assertEqual(plant.get(url).status_code, 403, url)
        for url in ("/", "/monitor", "/review", "/alerts", "/reports"):
            self.assertEqual(viewer.get(url).status_code, 200, url)
            self.assertEqual(plant.get(url).status_code, 200, url)
        self.assertEqual(viewer.get("/upload").status_code, 403)                     # head-office viewers do not upload
        self.assertEqual(plant.get("/upload").status_code, 200)
        self.assertEqual(admin.get("/upload").status_code, 200)

    def test_a_plant_user_only_ever_sees_their_own_plants(self):
        c = self.login("plant1")
        self.assertEqual(c.get("/plant/PU1").status_code, 200)
        for other in ("NA1", "KT1"):
            self.assertEqual(c.get(f"/plant/{other}").status_code, 403, other)
        body = c.get("/monitor").get_data(as_text=True)
        self.assertEqual(set(re.findall(r'/plant/([A-Z0-9]+)"', body)), {"PU1"})
        self.assertNotIn("NA1", c.get("/monitor?plant=NA1").get_data(as_text=True).split("<tbody>")[1].split("</tbody>")[0])
        self.assertEqual(c.get(f"/e/{self.exts['na']}").status_code, 403)             # another plant's bill viewer
        self.assertEqual(c.get(f"/e/{self.exts['na']}/file").status_code, 403)          # ... and its original document
        self.assertEqual(c.get(f"/e/{self.exts['na']}/page/1.png").status_code, 403)
        self.assertEqual(c.get(f"/e/{self.exts['pu']}").status_code, 200)
        self.assertNotIn("Nashik", c.get("/alerts?status=ALL").get_data(as_text=True))
        wb = load_workbook(io.BytesIO(c.get("/reports/register.xlsx").data))
        plants = {row[0] for row in wb["Report"].iter_rows(min_row=7, max_col=1, values_only=True) if row[0]}
        self.assertEqual(plants, {"PU1"})
        self.assertEqual(self.login("noplants").get("/monitor").get_data(as_text=True).count("/plant/"), 0)       # a regional user with no plants sees none

    def test_the_oracle_address_is_shown_to_administrators_only(self):
        adm = self.login("admin").get("/api/oracle-status").get_json()
        plant = self.login("plant1").get("/api/oracle-status").get_json()
        self.assertNotIn(":", plant["label"])                                              # no host:port for ordinary users
        self.assertEqual(adm["state"], plant["state"])

    def test_the_dashboard_reflects_stored_data_and_says_what_is_unverified(self):
        html = self.login("admin").get("/").get_data(as_text=True)
        self.assertIn("3 of 3 bills", html)
        self.assertIn("not been checked by a person", html)
        self.assertIn("ACTUAL", html)
        self.assertIn("THEORETICAL", html)
        self.assertIn("5,000", html)                                                  # the PF penalty printed on the PU1 bill
        verified = self.login("admin").get("/?quality=verified").get_data(as_text=True)
        self.assertIn("No bills in this selection", verified)

    def test_monitor_filters_sorting_and_search(self):
        c = self.login("admin")
        def plants(url):
            return re.findall(r'/plant/([A-Z0-9]+)"', c.get(url).get_data(as_text=True))
        self.assertEqual(sorted(plants("/monitor")), ["KT1", "NA1", "PU1"])
        self.assertEqual(plants("/monitor?plant=KT1"), ["KT1"])
        self.assertEqual(plants("/monitor?state=Kerala"), ["KT1"])
        self.assertEqual(plants("/monitor?exception=pf_critical"), ["NA1"])
        self.assertEqual(plants("/monitor?q=W-0001"), ["PU1"])
        self.assertEqual(plants("/monitor?sort=pf&dir=asc")[0], "NA1")
        self.assertEqual(plants("/monitor?sort=pf&dir=desc")[-1], "NA1")
        self.assertEqual(plants("/monitor?sort=evil;DROP&dir=sideways"), plants("/monitor"))          # sort keys are whitelisted

    def test_sql_in_a_filter_is_text(self):
        c = self.login("admin")
        for evil in ("' OR 1=1 --", "x'; DROP TABLE electricity_bills; --", "%' UNION SELECT password_hash FROM users --"):
            r = c.get("/monitor", query_string={"q": evil, "plant": evil, "state": evil})
            self.assertEqual(r.status_code, 200)
            self.assertEqual(len(re.findall(r'/plant/([A-Z0-9]+)"', r.get_data(as_text=True))), 0)
        self.assertEqual(len(self.rows("SELECT id FROM electricity_bills")), 3)

    def test_html_in_bill_data_is_escaped(self):
        body = self.login("admin").get("/rules").get_data(as_text=True) + self.login("admin").get("/discovery").get_data(as_text=True)
        self.assertNotIn("<script>alert(1)</script>", body)
        self.assertIn("&lt;script&gt;alert(1)&lt;/script&gt;", body)

    def test_no_secret_is_ever_sent_to_the_browser(self):
        c = self.login("admin")
        for url in self.PAGES + [f"/e/{self.exts['pu']}", "/reports/register.xlsx"]:
            r = c.get(url)
            blob = r.get_data(as_text=False)
            self.assertNotIn(ORACLE_SECRET.encode(), blob, url)
            self.assertNotIn(SMTP_SECRET.encode(), blob, url)
            self.assertNotIn(b"password_hash", blob, url)
        self.assertIn("&lt;set&gt;", c.get("/settings").get_data(as_text=True))          # only "set / empty" is ever shown

    def test_the_original_bill_is_served_unchanged_with_safe_headers(self):
        c = self.login("admin")
        r = c.get(f"/e/{self.exts['pu']}/file")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.mimetype, "application/pdf")
        self.assertEqual(r.data, (config.ORIGINALS_DIR / "t" / "w501.pdf").read_bytes())
        self.assertEqual(r.headers["X-Content-Type-Options"], "nosniff")
        self.assertIn("inline", r.headers.get("Content-Disposition", "inline"))

    def test_evidence_image_highlights_the_value(self):
        c = self.login("admin")
        r = c.get(f"/e/{self.exts['pu']}/page/1.png?field=kwh")
        self.assertEqual((r.status_code, r.mimetype), (200, "image/png"))
        self.assertTrue(r.data.startswith(b"\x89PNG"))
        plain = c.get(f"/e/{self.exts['pu']}/page/1.png")
        self.assertNotEqual(r.data, plain.data)                                            # the highlighted page differs from the plain one
        self.assertEqual(c.get(f"/e/{self.exts['pu']}/page/99.png").status_code, 404)

    def test_a_stored_path_outside_the_document_store_is_never_served(self):
        with storage.session(self.db) as conn:
            conn.execute("INSERT INTO electricity_attachments (attached_document_id, po_header_id, entity_name, pk1_value, document_id, file_name, fetch_status, sha256, local_path,"
                         " first_seen, last_seen) VALUES (999, 1, 'PO_HEADERS', '1', 999, 'evil.pdf', 'DOWNLOADED', 'shaevil', '../ops.db', 'x', 'x')")
            ext = conn.execute("INSERT INTO electricity_bill_extractions (attachment_id, sha256, reader_version, status, bill_status, created_at) "
                               "VALUES (999, 'shaevil', 'r', 'FAILED', 'FAILED', 'x')").lastrowid
        try:
            self.assertEqual(self.login("admin").get(f"/e/{ext}/file").status_code, 404)
        finally:
            with storage.session(self.db) as conn:
                conn.execute("DELETE FROM electricity_bill_extractions WHERE id = ?", (ext,))
                conn.execute("DELETE FROM electricity_attachments WHERE attached_document_id = 999")

    def test_excel_reports_are_real_workbooks_with_headers_formats_and_notes(self):
        c = self.login("admin")
        for kind in ("monitor", "consumption", "pf", "pf_penalties", "demand_violations", "demand_penalties", "monthly_cost", "savings", "register", "states", "exceptions"):
            r = c.get(f"/reports/{kind}.xlsx")
            self.assertEqual(r.status_code, 200, kind)
            self.assertIn("spreadsheetml", r.mimetype)
            wb = load_workbook(io.BytesIO(r.data))
            self.assertIn("Notes", wb.sheetnames)
            ws = wb["Report"]
            self.assertTrue(ws["A1"].value)
            self.assertEqual(ws.freeze_panes is not None, True)
            self.assertIn("Generated", ws["A3"].value)
        ws = load_workbook(io.BytesIO(c.get("/reports/pf_penalties.xlsx").data))["Report"]
        values = [r for r in ws.iter_rows(min_row=7, values_only=True) if r[0] == "PU1"]
        self.assertEqual(len(values), 1)
        headers = [c_.value for c_ in ws[6]]
        pen = values[0][headers.index("PF penalty (actual, on bill)")]
        self.assertEqual(pen, 5000.0)
        self.assertIsInstance(pen, (int, float))                                             # a real number, not text
        notes = " ".join(str(r[0]) for r in load_workbook(io.BytesIO(c.get("/reports/savings.xlsx").data))["Notes"].iter_rows(values_only=True) if r[0])
        self.assertIn("never added together", notes)
        self.assertEqual(c.get("/reports/not-a-report.xlsx").status_code, 404)

    def test_formula_text_from_a_bill_cannot_run_in_excel(self):
        with storage.session(self.db) as conn:
            conn.execute("UPDATE electricity_bills SET bill_number = '=HYPERLINK(\"http://evil\",\"x\")' WHERE extraction_id = ?", (self.exts["pu"],))
        try:
            ws = load_workbook(io.BytesIO(self.login("admin").get("/reports/register.xlsx").data))["Report"]
            cells = [c_ for row in ws.iter_rows(min_row=7) for c_ in row if isinstance(c_.value, str) and "HYPERLINK" in c_.value]
            self.assertTrue(cells)
            self.assertTrue(all(c_.data_type == "s" for c_ in cells))
        finally:
            with storage.session(self.db) as conn:
                conn.execute("UPDATE electricity_bills SET bill_number = 'W-0001' WHERE extraction_id = ?", (self.exts["pu"],))

    def test_state_changing_requests_need_the_csrf_token(self):
        c = self.login("admin")
        for url in (f"/e/{self.exts['pu']}/correct", f"/e/{self.exts['pu']}/verify", "/alerts/1/ack", "/automation/scan-now", "/automation/retry-failed", "/tariffs/1/approve", "/settings"):
            r = c.post(url, data={"field": "kwh", "value": "1"})
            self.assertEqual(r.status_code, 400, url)
        self.assertEqual(c.post("/logout").status_code, 400)
        self.assertEqual(c.post("/api/bg-settings", json={"manual_theme": "Night"}).status_code, 400)
        self.assertEqual(c.get("/logout").status_code, 405)                                  # logging out is not a link anybody can plant


# ------------------------------------------------------------------------------------------------ actions that change data
class ActionWebTests(WebBase):
    def setUp(self):
        self.build()
        self.addCleanup(self.destroy)

    def test_correcting_a_value_in_the_browser_updates_the_bill_keeps_the_original_and_audits(self):
        c = self.login("admin")
        ext = self.exts["pu"]
        r = self.post(c, f"/e/{ext}/correct", {"field": "kwh", "value": "91,000", "reason": "OCR dropped a digit"})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        self.assertEqual(self.rows("SELECT kwh FROM electricity_bills WHERE extraction_id = ?", ext)[0]["kwh"], 91000.0)
        f = self.rows("SELECT * FROM electricity_bill_fields WHERE extraction_id = ? AND field_name = 'kwh'", ext)[0]
        self.assertEqual((f["verification_status"], f["extracted_value_text"], f["corrected_by"]), ("CORRECTED", "90000.0", "admin"))
        self.assertEqual(self.rows("SELECT action, username FROM electricity_audit_logs WHERE action = 'CORRECT_FIELD'")[0]["username"], "admin")
        html = c.get(f"/e/{ext}").get_data(as_text=True)
        self.assertIn("originally read as 90000.0", html)
        self.assertIn("OCR dropped a digit", html)

    def test_bad_corrections_are_refused_with_a_message(self):
        c = self.login("admin")
        ext = self.exts["pu"]
        self.assertEqual(self.post(c, f"/e/{ext}/correct", {"field": "kwh", "value": "lots"}).status_code, 400)
        self.assertEqual(self.post(c, f"/e/{ext}/correct", {"field": "bill_date", "value": "31/02/2026"}).status_code, 400)
        r = self.post(c, f"/e/{ext}/correct", {"field": "password_hash; DROP", "value": "1"})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(self.rows("SELECT id FROM electricity_manual_corrections"), [])

    def test_only_an_administrator_corrects_oracle_bills(self):
        for u in ("viewer", "plant1"):
            c = self.login(u)
            self.assertEqual(self.post(c, f"/e/{self.exts['pu']}/correct", {"field": "kwh", "value": "1"}).status_code, 403, u)
            self.assertEqual(self.post(c, f"/e/{self.exts['pu']}/retry").status_code, 403, u)
            self.assertEqual(self.post(c, f"/e/{self.exts['pu']}/verify", {"field": "kwh"}).status_code, 403, u)
        self.assertEqual(self.rows("SELECT id FROM electricity_manual_corrections"), [])

    def test_verify_marks_fields_checked_and_the_bill_becomes_verified(self):
        c = self.login("admin")
        ext = self.exts["kt"]
        flagged = [r["field_name"] for r in self.rows("SELECT field_name FROM electricity_bill_fields WHERE extraction_id = ? AND verification_status = 'NEEDS_REVIEW'", ext)]
        self.assertTrue(flagged)
        r = c.post(f"/e/{ext}/verify", data={"field": flagged}, headers={"X-CSRF-Token": self.token(c)})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self.rows("SELECT bill_status FROM electricity_bill_extractions WHERE id = ?", ext)[0]["bill_status"], "VERIFIED")
        self.assertEqual(self.rows("SELECT quality FROM electricity_bills WHERE extraction_id = ?", ext)[0]["quality"], "VERIFIED")
        self.assertIn("Verified by a person only", c.get("/?quality=verified").get_data(as_text=True))
        self.assertNotIn("No bills in this selection", c.get("/?quality=verified").get_data(as_text=True))

    def test_a_bill_that_needs_review_cannot_be_closed_as_resolved(self):
        c = self.login("admin")
        r = self.post(c, f"/e/{self.exts['pu']}/resolve", {"note": "ignore it"})
        self.assertEqual(r.status_code, 400)
        self.assertIn("verifying or correcting", r.get_json()["error"])

    def test_alert_acknowledge_and_resolve_through_the_browser(self):
        c = self.login("admin")
        aid = self.rows("SELECT id FROM electricity_alerts WHERE alert_type = 'PF_PENALTY_ACTUAL'")[0]["id"]
        self.assertEqual(self.post(c, f"/alerts/{aid}/ack").status_code, 200)
        self.assertEqual(self.rows("SELECT status, acknowledged_by FROM electricity_alerts WHERE id = ?", aid)[0]["acknowledged_by"], "admin")
        self.assertEqual(self.post(c, f"/alerts/{aid}/resolve", {"note": "disputed"}).status_code, 200)
        self.assertEqual(self.rows("SELECT status FROM electricity_alerts WHERE id = ?", aid)[0]["status"], "RESOLVED")
        plant = self.login("plant2")                                                        # NA1's user must not touch PU1's alert
        self.assertEqual(self.post(plant, f"/alerts/{aid}/ack").status_code, 403)

    def test_tariff_rule_lifecycle_and_its_effect_on_the_analysis(self):
        c = self.login("admin")
        bad = c.post("/tariffs", data={"csrf_token": self.token(c), "state": "Maharashtra", "effective_from": "2026-04-01", "pf_penalty_rule_json": "{not json"}, follow_redirects=True)
        self.assertIn("not valid JSON", bad.get_data(as_text=True))
        nosrc = c.post("/tariffs", data={"csrf_token": self.token(c), "state": "Maharashtra", "effective_from": "2026-04-01", "contract_demand_tolerance_pct": "5"}, follow_redirects=True)
        self.assertIn("Draft rule #1 saved", nosrc.get_data(as_text=True))
        r = self.post(c, "/tariffs/1/approve")
        self.assertEqual(r.status_code, 400)                                                  # no source reference => cannot be approved
        self.assertIn("source", r.get_json()["error"])
        c.post("/tariffs", data={"csrf_token": self.token(c), "state": "Maharashtra", "effective_from": "2026-04-01", "contract_demand_tolerance_pct": "5",
                                 "source_reference": "MERC order 2026/04"}, follow_redirects=True)
        self.assertEqual(self.rows("SELECT d.expected_status FROM electricity_demand_analysis d JOIN electricity_bills b ON b.id = d.bill_id WHERE b.plant_code = 'PU1'")[0]["expected_status"], "UNAVAILABLE_NO_RULE")
        self.assertEqual(self.post(c, "/tariffs/2/approve").status_code, 200)
        d = self.rows("SELECT d.violation_status, d.tariff_rule_id FROM electricity_demand_analysis d JOIN electricity_bills b ON b.id = d.bill_id WHERE b.plant_code = 'PU1'")[0]
        self.assertEqual(d["tariff_rule_id"], 2)
        self.assertEqual(d["violation_status"], "WITHIN_PERMITTED_LIMIT")                       # MD 92 of CD 100, tolerance 5 %
        self.assertEqual(self.post(c, "/tariffs/2/draft").status_code, 200)                    # withdrawing it removes its effect again
        self.assertIsNone(self.rows("SELECT d.tariff_rule_id FROM electricity_demand_analysis d JOIN electricity_bills b ON b.id = d.bill_id WHERE b.plant_code = 'PU1'")[0]["tariff_rule_id"])
        self.assertEqual(len(self.rows("SELECT id FROM electricity_audit_logs WHERE action LIKE 'TARIFF_RULE_%'")), 4)

    def test_settings_validate_audit_and_reanalyse(self):
        c = self.login("admin")
        r = c.post("/settings", data={"csrf_token": self.token(c), "pf_excellent": "0.99", "pf_good": "0.85", "pf_attention": "0.90"}, follow_redirects=True)
        self.assertIn("PF bands must satisfy", r.get_data(as_text=True))
        form = {"csrf_token": self.token(c), "pf_good": "0.80", "pf_attention": "0.70", "alert_email_to": "ops@rdc.in", "alert_email_enabled": "on", "scan_interval_minutes": "120"}
        r = c.post("/settings", data=form, follow_redirects=True)
        self.assertIn("Settings saved", r.get_data(as_text=True))
        self.assertEqual(self.rows("SELECT performance_status FROM electricity_pf_analysis p JOIN electricity_bills b ON b.id = p.bill_id WHERE b.plant_code = 'NA1'")[0]["performance_status"], "CRITICAL")
        self.assertEqual(self.rows("SELECT performance_status FROM electricity_pf_analysis p JOIN electricity_bills b ON b.id = p.bill_id WHERE b.plant_code = 'PU1'")[0]["performance_status"], "GOOD")   # PF 0.90 >= new good 0.80
        log = self.rows("SELECT details_json FROM electricity_audit_logs WHERE action = 'SETTINGS_CHANGE'")[0]["details_json"]
        self.assertIn("pf_good", log)
        c.post("/settings", data={**form, "alert_email_to": "not-an-address"})
        self.assertIn("ops@rdc.in", self.rows("SELECT value FROM electricity_state WHERE key = 'cfg.alert_email_to'")[0]["value"])

    def test_identification_rules_can_be_changed_and_are_re_applied_without_oracle(self):
        c = self.login("admin")
        c.post("/rules", data={"csrf_token": self.token(c), "op": "add", "rule_type": "SUPPLIER_EXCLUDE", "value": "NASHIK", "note": "test"}, follow_redirects=True)
        self.assertEqual(self.rows("SELECT eligible FROM electricity_pos WHERE po_header_id = 2")[0]["eligible"], 0)
        self.assertNotIn("NA1", re.findall(r'/plant/([A-Z0-9]+)"', c.get("/monitor").get_data(as_text=True)))        # its bill leaves the dashboards ...
        self.assertEqual(len(self.rows("SELECT id FROM electricity_bills WHERE plant_code = 'NA1'")), 1)                 # ... but nothing is deleted
        rid = self.rows("SELECT id FROM electricity_id_rules WHERE value = 'NASHIK'")[0]["id"]
        c.post("/rules", data={"csrf_token": self.token(c), "op": "off", "rule_id": rid}, follow_redirects=True)
        self.assertEqual(self.rows("SELECT eligible FROM electricity_pos WHERE po_header_id = 2")[0]["eligible"], 1)
        item = self.rows("SELECT id FROM electricity_id_rules WHERE rule_type = 'ITEM_CODE'")[0]["id"]
        r = c.post("/rules", data={"csrf_token": self.token(c), "op": "off", "rule_id": item}, follow_redirects=True)
        self.assertIn("at least one active ITEM_CODE", r.get_data(as_text=True))

    def test_scan_now_without_an_in_process_scheduler_asks_the_worker_and_retry_failed_requeues(self):
        c = self.login("admin")
        with storage.session(self.db) as conn:
            conn.execute("UPDATE electricity_attachments SET fetch_status = 'FAILED', retry_count = 5 WHERE attached_document_id = 503")
        r = self.post(c, "/automation/scan-now")
        self.assertEqual(r.status_code, 200)
        r = self.post(c, "/automation/retry-failed")
        self.assertEqual(r.status_code, 200)
        a = self.rows("SELECT fetch_status, retry_count FROM electricity_attachments WHERE attached_document_id = 503")[0]
        self.assertEqual((a["fetch_status"], a["retry_count"]), ("PENDING", 0))
        self.assertEqual(len(self.rows("SELECT id FROM electricity_audit_logs WHERE action IN ('SCAN_NOW', 'RETRY_FAILED')")), 2)

    def test_retry_extraction_from_the_browser_queues_a_job_that_re_reads_the_bill(self):
        c = self.login("admin")
        self.assertEqual(self.post(c, f"/e/{self.exts['kt']}/retry").status_code, 200)
        self.assertEqual(len(self.rows("SELECT id FROM electricity_jobs WHERE kind = 'RETRY_EXTRACTION' AND status = 'QUEUED'")), 1)
        jobs.drain("t", self.db, run_ocr=False)
        self.assertEqual(self.rows("SELECT COUNT(*) AS n FROM electricity_bill_extractions WHERE archived = 1")[0]["n"], 1)
        self.assertEqual(len(self.rows("SELECT id FROM electricity_bills WHERE plant_code = 'KT1'")), 1)

    def test_billing_basis_can_be_set_from_the_bill_page(self):
        c = self.login("admin")
        self.assertEqual(self.post(c, f"/e/{self.exts['pu']}/basis", {"basis": "kVAh", "reason": "tariff says so"}).status_code, 200)
        self.assertEqual(self.rows("SELECT billing_basis FROM electricity_bills WHERE plant_code = 'PU1'")[0]["billing_basis"], "kVAh")
        self.assertEqual(self.post(c, f"/e/{self.exts['pu']}/basis", {"basis": "nonsense"}).status_code, 400)


# ------------------------------------------------------------------------------------------------ upload through the browser
class UploadWebTests(WebBase):
    def setUp(self):
        self.build()
        self.addCleanup(self.destroy)

    def upload(self, c, data, name="bill.pdf", **fields):
        form = {"csrf_token": self.token(c), "plant_code": "PU1", **fields, "file": (io.BytesIO(data), name)}
        return c.post("/upload", data=form, content_type="multipart/form-data")

    def good(self, **kw):
        d = tempfile.mkdtemp()
        return synth.make_bill(pathlib.Path(d) / "u.pdf", **kw).read_bytes()

    def test_upload_read_verify_flow_for_a_plant_user(self):
        c = self.login("plant1")
        r = self.upload(c, self.good(bill_no="UPL-0001", consumer="777777777777"), consumer_number="777777777777", billing_month="2026-09")
        self.assertEqual(r.status_code, 302)
        status_url = r.headers["Location"]
        self.assertIn("/upload/", status_url)
        self.assertEqual(c.get(status_url).status_code, 200)                                  # the "reading…" page
        uid = int(status_url.rsplit("/", 1)[1])
        self.assertEqual(c.get(f"/api/upload/{uid}").get_json()["status"], "QUEUED")
        jobs.drain("t", self.db, run_ocr=False)
        api = c.get(f"/api/upload/{uid}").get_json()
        self.assertEqual(api["status"], "PROCESSED")
        r = c.get(status_url)
        self.assertEqual(r.status_code, 302)                                                    # now it forwards to the editable bill form
        page = c.get(r.headers["Location"])
        html = page.get_data(as_text=True)
        self.assertIn("Correct a value", html)                        # the uploader may correct their own upload
        ext = api["extraction_id"]
        self.assertEqual(self.post(c, f"/e/{ext}/correct", {"field": "kwh", "value": "90500"}).status_code, 200)
        self.assertEqual(self.rows("SELECT kwh, source FROM electricity_bills WHERE extraction_id = ?", ext)[0], {"kwh": 90500.0, "source": "MANUAL_UPLOAD"})
        other = self.login("plant2")
        self.assertEqual(other.get(f"/e/{ext}").status_code, 403)                              # a different plant's user cannot open it

    def test_a_plant_user_cannot_upload_for_another_plant(self):
        c = self.login("plant1")
        r = self.upload(c, self.good(), plant_code="NA1")
        self.assertEqual(r.status_code, 403)
        self.assertEqual(self.rows("SELECT id FROM electricity_uploads"), [])

    def test_hostile_and_wrong_files_are_refused_with_a_message(self):
        c = self.login("plant1")
        for data, name, frag in ((b"MZ\x90\x00 an executable", "bill.pdf", "Only PDF or image"), (b"", "empty.pdf", "empty"),
                                 (b"%PDF-1.4\n<< /JavaScript (app.alert(1)) /S /JavaScript >>", "x.pdf", "active content")):
            r = self.upload(c, data, name)
            self.assertEqual(r.status_code, 200, name)
            self.assertIn(frag, r.get_data(as_text=True))
        self.assertEqual(self.rows("SELECT id FROM electricity_uploads"), [])
        self.assertEqual(self.upload(c, self.good(), plant_code="").status_code, 200)

    def test_a_duplicate_upload_shows_the_earlier_result(self):
        c = self.login("plant1")
        data = self.good(bill_no="UPL-0002", consumer="888888888888")
        first = self.upload(c, data)
        second = self.upload(c, data)
        self.assertEqual(first.headers["Location"], second.headers["Location"])
        self.assertEqual(len(self.rows("SELECT id FROM electricity_uploads")), 1)

    def test_the_upload_page_lists_only_the_users_plants(self):
        html = self.login("plant1").get("/upload").get_data(as_text=True)
        self.assertIn("PU1", html)
        self.assertNotIn("NA1 — Nashik", html)


if __name__ == "__main__":
    unittest.main()
